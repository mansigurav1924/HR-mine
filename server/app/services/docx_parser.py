import docx
import io
import re
import logging
from typing import List, Dict, Any

logger = logging.getLogger(__name__)

def extract_structured_docx(file_bytes: bytes) -> Dict[str, Any]:
    """
    Multi-strategy DOCX extraction using python-docx:
    1. Paragraph text
    2. Table cells across all document tables
    3. Section headers and footers
    4. Hyperlink relationship target URLs
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Failed to parse DOCX document: {str(e)}")

    paragraphs = []
    tables_data = []
    headers_footers = []
    links = []
    
    # 1. Paragraphs
    for para in doc.paragraphs:
        p_text = para.text.strip()
        if p_text:
            paragraphs.append(p_text)
            
    # 2. Tables
    for table in doc.tables:
        t_matrix = []
        for row in table.rows:
            row_cells = []
            for cell in row.cells:
                c_text = cell.text.strip()
                # Deduplicate identical adjacent cells (due to merged cells in word)
                if not row_cells or row_cells[-1] != c_text:
                    row_cells.append(c_text)
            if any(c for c in row_cells if c):
                t_matrix.append(row_cells)
        if t_matrix:
            tables_data.append(t_matrix)

    # 3. Section Headers & Footers
    try:
        for section in doc.sections:
            if section.header and section.header.paragraphs:
                for hp in section.header.paragraphs:
                    ht = hp.text.strip()
                    if ht and ht not in headers_footers:
                        headers_footers.append(ht)
            if section.footer and section.footer.paragraphs:
                for fp in section.footer.paragraphs:
                    ft = fp.text.strip()
                    if ft and ft not in headers_footers:
                        headers_footers.append(ft)
    except Exception as e:
        logger.debug(f"Error reading docx headers/footers: {e}")

    # 4. Hyperlink relationships
    try:
        for rel in doc.part.rels.values():
            if "hyperlink" in rel.reltype:
                target = rel.target_ref
                if target and isinstance(target, str) and target.strip():
                    cleaned_target = target.strip()
                    if cleaned_target not in links:
                        links.append(cleaned_target)
    except Exception as e:
        logger.debug(f"Error reading docx hyperlink rels: {e}")

    # Build unified plain text combining paragraphs, tables, and headers
    unified_parts = []
    if headers_footers:
        unified_parts.extend(headers_footers)
    unified_parts.extend(paragraphs)
    
    for t in tables_data:
        for row in t:
            unified_parts.append(" | ".join([c for c in row if c]))

    unified_text = "\n\n".join(unified_parts).strip()
    is_scanned = len(unified_text) < 30 and len(paragraphs) == 0 and len(tables_data) == 0

    return {
        "plain_text": unified_text,
        "paragraphs": paragraphs,
        "tables": tables_data,
        "headers_footers": headers_footers,
        "links": links,
        "is_scanned": is_scanned
    }

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Convenience backward-compatible text extractor."""
    data = extract_structured_docx(file_bytes)
    return data["plain_text"]

def extract_links_from_docx(file_bytes: bytes) -> List[str]:
    """Convenience backward-compatible link extractor."""
    data = extract_structured_docx(file_bytes)
    return data["links"]
