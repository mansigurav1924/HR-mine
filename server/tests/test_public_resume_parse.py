import io
import uuid
import fitz
import docx
import pytest
from fastapi.testclient import TestClient
from app.main import app
from app.api.public_applications import _PARSE_RATE_LIMITS
from app.services.supabase_client import supabase
from app.services.profile_extractor import clean_and_normalize_url, extract_urls

client = TestClient(app)

def create_sample_pdf(text: str) -> bytes:
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    return doc.write()

def create_pdf_with_embedded_links(text: str, links: list) -> bytes:
    """Creates a PDF with clickable URI annotations over text."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), text)
    y = 60
    for uri in links:
        rect = fitz.Rect(50, y, 200, y + 15)
        page.insert_link({"kind": fitz.LINK_URI, "from": rect, "uri": uri})
        y += 20
    return doc.write()

def create_sample_docx(paragraphs: list) -> bytes:
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def create_docx_with_hyperlinks(paragraphs: list, hyperlink_uris: list) -> bytes:
    """Creates a DOCX with external hyperlink relationships in part.rels."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    for uri in hyperlink_uris:
        doc.part.relate_to(uri, docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

@pytest.fixture(autouse=True)
def reset_rate_limits():
    _PARSE_RATE_LIMITS.clear()

def test_public_resume_parse_pdf_complete_fields():
    resume_text = """
Aarav Sharma
aarav.sharma2026@gmail.com | +91 98765 43210 | Bengaluru, Karnataka, India
GitHub: github.com/aaravsharma | LinkedIn: linkedin.com/in/aarav-sharma-dev | Portfolio: https://aarav.dev

EDUCATION
R.V. College of Engineering, Bengaluru
Bachelor of Engineering (B.E.) in Computer Science and Engineering
CGPA: 8.92 / 10.0 | Expected Graduation: 2026

TECHNICAL SKILLS
Languages: Python, JavaScript, TypeScript, Java, SQL, C++
Frameworks: React, Node.js, FastAPI, Express.js, Tailwind CSS
Databases: PostgreSQL, MongoDB, MySQL, Supabase
Tools: Git, Docker, Postman, Linux

WORK EXPERIENCE
Full Stack Intern — TechVanguard Solutions (Jan 2025 – Present)
• Developed responsive React UI components and FastAPI backend endpoints.
• Implemented JWT authentication and PostgreSQL database connection pooling.

PROJECTS
AI Resume Parser & Job Matcher (Python, FastAPI, Scikit-Learn)
• Built NLP resume extraction pipeline extracting candidate details and skills.
• Deployed microservice with Docker and automated testing.

Campus Event Management Portal (React, Node.js, MongoDB)
• Designed full-stack portal used by 2,000+ students for event registrations.
"""
    pdf_bytes = create_sample_pdf(resume_text)
    files = {"resume": ("aarav_resume.pdf", pdf_bytes, "application/pdf")}
    
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 200
    data = response.json()
    
    # 1. Verification of status and privacy
    assert data["parsed"] is True
    assert data["manual_entry_required"] is False
    assert "raw_text" not in data
    assert "classification" not in data
    
    # 2. Personal info extraction
    assert data["personal"]["full_name"] == "Aarav Sharma"
    assert data["personal"]["email"] == "aarav.sharma2026@gmail.com"
    assert "98765" in data["personal"]["phone"]
    assert "Bengaluru" in data["personal"]["location"]
    
    # 3. Education extraction
    assert "R.V. College of Engineering" in data["education"]["college"]
    assert data["education"]["degree"] == "B.E."
    assert "Computer Science" in data["education"]["specialization"]
    assert "2026" in data["education"]["graduation_year"]
    assert "8.92" in data["education"]["cgpa"]
    
    # 4. Skills extraction
    for s in ["Python", "React", "FastAPI", "SQL", "Docker", "PostgreSQL"]:
        assert s in data["skills"]
        
    # 5. Experience extraction
    assert "TechVanguard Solutions" in data["experience"]["summary"]
    assert len(data["experience"]["internships"]) > 0
    assert data["experience"]["internships"][0]["company"] == "TechVanguard Solutions"
    assert data["experience"]["internships"][0]["role"] == "Full Stack Intern"
    
    # 6. Projects extraction
    assert len(data["projects"]) > 0
    assert "AI Resume Parser" in data["projects"][0]["name"]
    assert "Campus Event Management Portal" in data["projects_summary"]
    
    # 7. Online links extraction
    assert data["links"]["github_url"] == "https://github.com/aaravsharma"
    assert data["links"]["linkedin_url"] == "https://linkedin.com/in/aarav-sharma-dev"
    assert data["links"]["portfolio_url"] == "https://aarav.dev"

def test_pdf_with_hidden_embedded_hyperlinks():
    """Tests PDF where 'GitHub', 'LinkedIn', 'Portfolio' text hides the actual URI in PDF annotations."""
    resume_text = """
Kiran Kumar
kiran@example.com | 9876543210
GitHub
LinkedIn
Portfolio
EDUCATION: NIT Surathkal, B.Tech Computer Science
SKILLS: React, Python, PostgreSQL
"""
    embedded_uris = [
        "https://github.com/kirankumar-dev",
        "https://www.linkedin.com/in/kiran-kumar-swe",
        "https://kirankumar.vercel.app"
    ]
    pdf_bytes = create_pdf_with_embedded_links(resume_text, embedded_uris)
    files = {"resume": ("kiran.pdf", pdf_bytes, "application/pdf")}
    
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 200
    data = response.json()
    
    assert data["parsed"] is True
    assert data["links"]["github_url"] == "https://github.com/kirankumar-dev"
    assert data["links"]["linkedin_url"] == "https://linkedin.com/in/kiran-kumar-swe"
    assert data["links"]["portfolio_url"] == "https://kirankumar.vercel.app"
    # Flat convenience fields also populated
    assert data["github_url"] == "https://github.com/kirankumar-dev"
    assert data["linkedin_url"] == "https://linkedin.com/in/kiran-kumar-swe"
    assert data["portfolio_url"] == "https://kirankumar.vercel.app"

def test_docx_with_hidden_embedded_hyperlinks():
    """Tests DOCX where hyperlink targets are stored in document relationships."""
    paragraphs = [
        "Siddharth Roy",
        "siddharth@example.com | 9876543211",
        "EDUCATION: IIT Delhi, B.Tech Computer Science",
        "SKILLS: Python, Machine Learning, Docker, SQL"
    ]
    hyperlinks = [
        "https://github.com/siddharthroy-ai",
        "https://linkedin.com/in/siddharth-roy",
        "https://siddharth-portfolio.netlify.app"
    ]
    docx_bytes = create_docx_with_hyperlinks(paragraphs, hyperlinks)
    files = {"resume": ("sid.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 200
    data = response.json()
    
    assert data["parsed"] is True
    assert data["links"]["github_url"] == "https://github.com/siddharthroy-ai"
    assert data["links"]["linkedin_url"] == "https://linkedin.com/in/siddharth-roy"
    assert data["links"]["portfolio_url"] == "https://siddharth-portfolio.netlify.app"

def test_url_normalization_and_punctuation_stripping():
    """Tests URLs without https:// and with trailing punctuation."""
    assert clean_and_normalize_url("github.com/alex123.") == "https://github.com/alex123"
    assert clean_and_normalize_url("www.linkedin.com/in/alex-smith/") == "https://www.linkedin.com/in/alex-smith/"
    assert clean_and_normalize_url("<https://alex.dev/portfolio>,") == "https://alex.dev/portfolio"

    text = """
Neha Gupta
neha@example.com | +91 9988776655
Profile Links: (github.com/nehagupta), [www.linkedin.com/in/neha-gupta], {https://neha.tech}
Skills: React, Node.js, Python
"""
    pdf_bytes = create_sample_pdf(text)
    files = {"resume": ("neha.pdf", pdf_bytes, "application/pdf")}
    
    res = client.post("/api/public/resume/parse", files=files)
    assert res.status_code == 200
    links = res.json()["links"]
    assert links["github_url"] == "https://github.com/nehagupta"
    assert links["linkedin_url"] == "https://linkedin.com/in/neha-gupta"
    assert links["portfolio_url"] == "https://neha.tech"

def test_duplicate_github_urls_deduplicated():
    """Ensures multiple mentions of GitHub are deduplicated cleanly."""
    text = """
Rohan Das
rohan@example.com
GitHub: https://github.com/rohandas
Project Repo: github.com/rohandas/cool-app
Other Repo: https://github.com/rohandas/ai-bot
Skills: Python, FastAPI
"""
    pdf_bytes = create_sample_pdf(text)
    files = {"resume": ("rohan.pdf", pdf_bytes, "application/pdf")}
    
    res = client.post("/api/public/resume/parse", files=files)
    assert res.status_code == 200
    links = res.json()["links"]
    assert links["github_url"] == "https://github.com/rohandas"

def test_missing_portfolio_returns_empty_cleanly():
    """Ensures when portfolio is missing, it returns empty string or None without crashing."""
    text = """
Ananya Roy
ananya@example.com | 9876543210
GitHub: github.com/ananyaroy
LinkedIn: linkedin.com/in/ananya-roy
Skills: Java, Spring Boot, MySQL
"""
    pdf_bytes = create_sample_pdf(text)
    files = {"resume": ("ananya.pdf", pdf_bytes, "application/pdf")}
    
    res = client.post("/api/public/resume/parse", files=files)
    assert res.status_code == 200
    data = res.json()
    assert data["links"]["portfolio_url"] in ("", None)
    assert data["portfolio_url"] in ("", None)
    assert data["links"]["github_url"] == "https://github.com/ananyaroy"
    assert data["links"]["linkedin_url"] == "https://linkedin.com/in/ananya-roy"

def test_public_resume_parse_scanned_or_empty_pdf():
    empty_pdf = create_sample_pdf("")
    files = {"resume": ("scanned_resume.pdf", empty_pdf, "application/pdf")}
    
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 200
    data = response.json()
    
    assert data["parsed"] is False
    assert data["manual_entry_required"] is True
    assert "could not extract text" in data["message"].lower() or "couldn't extract text" in data["message"].lower()
    assert data["skills"] == []

def test_public_resume_parse_invalid_extension():
    files = {"resume": ("test.txt", b"Hello World text file", "text/plain")}
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 400
    assert "Only .pdf and .docx" in response.json()["detail"]

def test_public_resume_parse_corrupt_signature():
    files = {"resume": ("fake.pdf", b"NOT_A_VALID_PDF_HEADER", "application/pdf")}
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 400
    assert "Corrupted or invalid PDF" in response.json()["detail"]

def test_public_resume_parse_oversized_file():
    oversized = b"%PDF-" + b"0" * (5 * 1024 * 1024 + 1024)
    files = {"resume": ("huge.pdf", oversized, "application/pdf")}
    response = client.post("/api/public/resume/parse", files=files)
    assert response.status_code == 400
    assert "exceeds maximum allowed limit" in response.json()["detail"]

def test_public_resume_parse_does_not_create_application():
    init_res = supabase.table("applications").select("application_id", count="exact").execute()
    init_count = init_res.count if hasattr(init_res, 'count') and init_res.count is not None else len(init_res.data or [])
    
    pdf_bytes = create_sample_pdf("Jane Doe Python React SQL Experience at TechCorp Projects AI Tool")
    files = {"resume": ("resume.pdf", pdf_bytes, "application/pdf")}
    
    res = client.post("/api/public/resume/parse", files=files)
    assert res.status_code == 200
    
    after_res = supabase.table("applications").select("application_id", count="exact").execute()
    after_count = after_res.count if hasattr(after_res, 'count') and after_res.count is not None else len(after_res.data or [])
    
    assert after_count == init_count

def test_public_submission_with_reviewed_and_edited_resume_data():
    jobs_res = client.get("/api/public/jobs")
    assert jobs_res.status_code == 200
    jobs = jobs_res.json()
    pos_id = jobs[0]["position_id"]
    
    raw_pdf = create_sample_pdf("Maya Sen\nmaya.sen@example.com\n+91 9123456789\nBITS Pilani\nB.E. Computer Science\nSkills: React, Python, Docker\ngithub.com/mayasen")
    parse_res = client.post("/api/public/resume/parse", files={"resume": ("maya.pdf", raw_pdf, "application/pdf")})
    assert parse_res.status_code == 200
    parsed_info = parse_res.json()
    
    # Candidate custom edits
    edited_name = parsed_info["personal"]["full_name"] + " (Edited)"
    edited_email = f"maya_{uuid.uuid4().hex[:6]}@example.com"
    edited_skills = parsed_info["skills"] + ["FastAPI", "TypeScript"]
    
    submit_res = client.post(
        "/api/public/applications",
        data={
            "position_id": pos_id,
            "full_name": edited_name,
            "email": edited_email,
            "phone": "+91 9999988888",
            "college": parsed_info["education"]["college"],
            "degree": parsed_info["education"]["degree"],
            "current_year": "2026",
            "skills": ", ".join(edited_skills),
            "experience": "Edited 6-month internship summary",
            "projects": "Edited major AI project",
            "github_url": parsed_info["links"]["github_url"],
            "linkedin_url": "https://linkedin.com/in/mayasen-candidate",
            "portfolio_url": "https://mayasen.me",
            "consent_given": "true"
        },
        files={
            "resume": ("maya.pdf", raw_pdf, "application/pdf")
        }
    )
    assert submit_res.status_code == 200
    sub_data = submit_res.json()
    assert sub_data["success"] is True
    app_id = sub_data["application_id"]
    
    db_rec = supabase.table("applications").select("*").eq("application_id", app_id).single().execute()
    assert db_rec.data["candidate_name"] == edited_name
    assert db_rec.data["email"] == edited_email
    assert "FastAPI" in db_rec.data["skills"]
    assert db_rec.data["current_status"] in ("application_received", "under_review", "ml_evaluated")
