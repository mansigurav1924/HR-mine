import io
import pytest
import docx
import fitz
from app.services.pdf_parser import extract_structured_pdf, extract_text_from_pdf, extract_links_from_pdf, is_scanned_or_empty_pdf
from app.services.docx_parser import extract_structured_docx, extract_text_from_docx, extract_links_from_docx
from app.services.profile_extractor import (
    extract_full_profile, extract_name_candidates, extract_email_candidates,
    extract_phone_candidates, extract_location_candidates, extract_education_details,
    extract_skills_matrix, extract_urls, extract_experience_structured,
    extract_projects_structured, segment_sections, normalize_text
)

# ==============================================================================
# 1. TEXT NORMALIZATION & SECTION SEGMENTATION TESTS
# ==============================================================================

def test_text_normalization():
    raw = "John Doe \u2022 Developer \u2013 2024 \u00a0\u202f City, India"
    norm = normalize_text(raw)
    assert "•" in norm
    assert "-" in norm
    assert "\u00a0" not in norm
    assert "John Doe" in norm

def test_unusual_section_headings_segmentation():
    sample_resume = """
Johnathan Vance
john.vance@example.com | +91 9876543210 | Bengaluru, India

ACADEMIC PROFILE
B.Tech in Computer Science and Engineering
Indian Institute of Technology, Madras
CGPA: 8.9/10 (Batch of 2025)

CORE COMPETENCIES
Python, React, FastApi, PostgreSQL, Machine Learning, Docker, Git

PRACTICAL EXPERIENCE
Software Engineer Intern at TechCorp Solutions (Jan 2024 - Jun 2024)
• Built high-performance backend microservices using FastAPI and PostgreSQL.

KEY PROJECTS
Smart Recruitment AI (Python, React)
• Engineered automated resume parsing engine with hybrid matrix algorithms.

ONLINE PROFILES
GitHub: https://github.com/jvance
LinkedIn: https://linkedin.com/in/jvance
Portfolio: https://jvance.dev
"""
    sections = segment_sections(sample_resume)
    assert len(sections["EDUCATION"]) > 0
    assert len(sections["SKILLS"]) > 0
    assert len(sections["EXPERIENCE"]) > 0
    assert len(sections["PROJECTS"]) > 0

# ==============================================================================
# 2. INDIVIDUAL FIELD EXTRACTION TESTS
# ==============================================================================

def test_name_extraction():
    sample = """
Aarav Sharma
aarav.sharma@example.com | +91 9123456789
Bengaluru, Karnataka, India
"""
    sections = segment_sections(sample)
    name, conf = extract_name_candidates(sample, sections)
    assert name == "Aarav Sharma"
    assert conf >= 90

def test_email_and_mailto_extraction():
    # Regular text email
    sample = "Contact: Priya Patel (priya.patel2025@gmail.com)"
    email, conf = extract_email_candidates(sample)
    assert email == "priya.patel2025@gmail.com"
    assert conf >= 95

    # Mailto hyperlink candidate
    links = ["mailto:priya.patel2025@gmail.com?subject=Job"]
    email_link, conf_link = extract_email_candidates("", embedded_links=links)
    assert email_link == "priya.patel2025@gmail.com"
    assert conf_link == 100

def test_phone_formats_indian_and_international():
    # Indian with +91 and spaces
    p1, conf1 = extract_phone_candidates("Phone: +91 98765 43210")
    assert "98765" in p1
    assert conf1 >= 95

    # Indian 10 digits
    p2, conf2 = extract_phone_candidates("Mobile: 9876543210")
    assert "9876543210" in p2
    assert conf2 >= 95

    # Non-phone (years, CGPA, roll number) must not be extracted
    p_fake, _ = extract_phone_candidates("Year: 2020-2024, CGPA: 8.5/10, Roll: 20CS045")
    assert p_fake is None

def test_location_extraction():
    sample = """
Rohan Verma
rohan.v@example.com | +91 9876543210 | Hyderabad, India
"""
    sections = segment_sections(sample)
    loc, conf = extract_location_candidates(sections, sample)
    assert "Hyderabad" in loc
    assert conf >= 90

def test_education_degree_specialization_cgpa_and_year():
    sample = """
EDUCATION
Bachelor of Technology in Computer Science & Engineering
National Institute of Technology, Trichy
CGPA: 8.75 / 10 | Expected Graduation: 2025
"""
    sections = segment_sections(sample)
    edu, conf = extract_education_details(sections, sample)
    assert edu["college"] is not None and "National Institute of Technology" in edu["college"]
    assert edu["degree"] == "Bachelor's Degree" or edu["degree"] == "B.Tech"
    assert "Computer Science" in (edu["specialization"] or "")
    assert "8.75" in (edu["cgpa"] or "")
    assert edu["graduation_year"] == "2025"
    assert conf >= 95

def test_skills_extraction_and_domain_aliases():
    sample = """
SKILLS & TOOLS
Programming: Python, JS, TS, Java, C++
Frameworks: React.js, NodeJS, Express, FastAPI, Django, SpringBoot
Databases & Cloud: PostgreSQL, MongoDB, AWS, Docker, Kubernetes, Git
Data & ML: Scikit-learn, PyTorch, Pandas, NumPy, PowerBI, Excel
"""
    sections = segment_sections(sample)
    skills, conf = extract_skills_matrix(sections, sample)
    
    # Verify canonical alias mapping
    assert "JavaScript" in skills  # from JS
    assert "TypeScript" in skills  # from TS
    assert "React" in skills       # from React.js
    assert "Node.js" in skills     # from NodeJS
    assert "FastAPI" in skills
    assert "PostgreSQL" in skills
    assert "Scikit-learn" in skills # from Scikit-learn
    assert "PyTorch" in skills
    assert "Power BI" in skills    # from PowerBI
    assert conf == 100

def test_experience_and_internship_extraction():
    sample = """
WORK EXPERIENCE
Full Stack Intern - ABC Digital Labs (Jun 2024 - Aug 2024)
• Developed responsive frontend UI with React and Tailwind CSS.
• Integrated REST APIs with FastAPI backend.

Backend Developer Intern at CloudNine Tech (Jan 2024 - Apr 2024)
• Optimized database queries in PostgreSQL.
"""
    sections = segment_sections(sample)
    exp, conf = extract_experience_structured(sections, sample)
    assert len(exp["internships"]) >= 2
    assert exp["internships"][0]["role"] is not None
    assert "ABC Digital Labs" in exp["internships"][0]["company"] or "Full Stack Intern" in exp["internships"][0]["title"]
    assert conf >= 95

def test_projects_extraction():
    sample = """
PROJECTS
AI Resume Screener (Python, FastAPI, Scikit-learn)
• Developed automated resume classification model with 92% precision.
• Built interactive candidate review dashboard in React.

E-Commerce Web Platform (React, Node.js, MongoDB)
• Implemented full shopping cart and payment gateway integration.
"""
    sections = segment_sections(sample)
    projects, summary, conf = extract_projects_structured(sections, sample)
    assert len(projects) >= 2
    assert "AI Resume Screener" in projects[0]["name"]
    assert "Python" in projects[0]["technologies"] or "Scikit-learn" in projects[0]["technologies"]
    assert conf >= 95

def test_url_classification_github_linkedin_portfolio():
    sample = """
Profiles:
GitHub: https://github.com/developer-hero
LinkedIn: https://linkedin.com/in/developer-hero-2025
Portfolio: https://developer-hero.vercel.app
Docs: https://react.dev
Search: https://google.com
"""
    urls, confs = extract_urls(sample)
    assert urls["github_url"] == "https://github.com/developer-hero"
    assert urls["linkedin_url"] == "https://linkedin.com/in/developer-hero-2025"
    assert urls["portfolio_url"] == "https://developer-hero.vercel.app"
    # General sites should NOT be classified as portfolio
    assert "google.com" not in (urls["portfolio_url"] or "")
    assert "react.dev" not in (urls["portfolio_url"] or "")

def test_parser_never_invents_missing_data():
    empty_sample = "Just a short note with no structure."
    profile = extract_full_profile(empty_sample)
    assert profile["personal"]["email"] == ""
    assert profile["personal"]["phone"] == ""
    assert profile["links"]["github_url"] == ""
    assert profile["links"]["linkedin_url"] == ""
    assert profile["links"]["portfolio_url"] == ""
    assert profile["education"]["college"] == ""
    assert profile["skills"] == []

# ==============================================================================
# 3. PDF MULTI-STRATEGY & LAYOUT TESTS (PyMuPDF)
# ==============================================================================

def test_pdf_multi_strategy_and_embedded_links():
    doc = fitz.open()
    page = doc.new_page()
    
    # Add text
    page.insert_text((50, 50), "Dev Sharma\nFull Stack Developer\ndev.sharma@example.com\n+91 9876543210\nBengaluru, India")
    page.insert_text((50, 150), "TECHNICAL SKILLS\nReact, JavaScript, Python, FastAPI, PostgreSQL")
    page.insert_text((50, 250), "EDUCATION\nB.Tech Computer Science\nIndian Institute of Technology Delhi\nCGPA: 9.1 | 2025")
    
    # Add clickable link annotation hiding URL
    rect = fitz.Rect(50, 320, 150, 340)
    page.insert_text((50, 335), "GitHub Profile")
    page.insert_link({
        "kind": fitz.LINK_URI,
        "from": rect,
        "uri": "https://github.com/devsharma-code"
    })
    
    pdf_bytes = doc.tobytes()
    doc.close()

    # Extract structured PDF
    res = extract_structured_pdf(pdf_bytes)
    assert len(res["blocks"]) > 0
    assert len(res["words"]) > 0
    assert "https://github.com/devsharma-code" in res["links"]
    assert res["is_scanned"] is False

    # Extract full profile using Hybrid Matrix
    profile = extract_full_profile(res["plain_text"], embedded_links=res["links"], blocks=res["blocks"])
    assert profile["personal"]["full_name"] == "Dev Sharma"
    assert profile["personal"]["email"] == "dev.sharma@example.com"
    assert "9876543210" in profile["personal"]["phone"]
    assert profile["links"]["github_url"] == "https://github.com/devsharma-code"
    assert "React" in profile["skills"]
    assert "Python" in profile["skills"]
    assert profile["extraction_meta"]["confidence"]["email"] >= 95

def test_pdf_two_column_layout_handling():
    doc = fitz.open()
    page = doc.new_page()
    
    # Column 1 (Left Sidebar)
    page.insert_text((50, 50), "Vikram Singh\nvikram.s@example.com\n+91 9988776655\nPune, India\n\nSKILLS\nPython, Django, SQL")
    
    # Column 2 (Main Column)
    page.insert_text((300, 50), "EDUCATION\nB.E. Information Technology\nPune Institute of Computer Technology\nGraduation: 2024\n\nEXPERIENCE\nBackend Intern at PuneTech Labs\nBuilt APIs using Python and Django.")
    
    pdf_bytes = doc.tobytes()
    doc.close()

    res = extract_structured_pdf(pdf_bytes)
    profile = extract_full_profile(res["plain_text"], blocks=res["blocks"])
    assert profile["personal"]["full_name"] == "Vikram Singh"
    assert profile["personal"]["email"] == "vikram.s@example.com"
    assert "Python" in profile["skills"]
    assert "Pune Institute of Computer Technology" in (profile["education"]["college"] or "")

# ==============================================================================
# 4. DOCX MULTI-STRATEGY & TABLE EXTRACTION TESTS
# ==============================================================================

def test_docx_table_and_hyperlink_extraction():
    doc = docx.Document()
    doc.add_paragraph("Ananya Roy")
    doc.add_paragraph("ananya.roy@example.com | +91 9876501234 | Kolkata, India")
    
    # Add Education Table
    table = doc.add_table(rows=2, cols=4)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Degree"
    hdr_cells[1].text = "Institution"
    hdr_cells[2].text = "Year"
    hdr_cells[3].text = "CGPA"
    
    row_cells = table.rows[1].cells
    row_cells[0].text = "BCA"
    row_cells[1].text = "St. Xaviers College, Kolkata"
    row_cells[2].text = "2025"
    row_cells[3].text = "8.8"

    doc.add_paragraph("SKILLS\nReact, JavaScript, HTML, CSS, Tailwind CSS, SQL")

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    res = extract_structured_docx(docx_bytes)
    assert len(res["tables"]) >= 1
    assert "St. Xaviers College, Kolkata" in res["plain_text"]

    profile = extract_full_profile(res["plain_text"], tables=res["tables"])
    assert profile["personal"]["full_name"] == "Ananya Roy"
    assert profile["personal"]["email"] == "ananya.roy@example.com"
    assert "St. Xaviers College" in (profile["education"]["college"] or "")
    assert profile["education"]["degree"] == "BCA"
    assert "React" in profile["skills"]
    assert "JavaScript" in profile["skills"]

def test_scanned_pdf_heuristic():
    assert is_scanned_or_empty_pdf("") is True
    assert is_scanned_or_empty_pdf("   ") is True
    assert is_scanned_or_empty_pdf("Short") is True
    assert is_scanned_or_empty_pdf("This is a full resume with lots of text content over thirty characters.") is False
