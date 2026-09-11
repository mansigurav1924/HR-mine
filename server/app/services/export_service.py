import io
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from typing import Optional, Dict, Any, List, Tuple
from datetime import datetime
from app.services.supabase_client import supabase

MAX_EXPORT_LIMIT = 5000

class ExportService:
    def export_resource(
        self,
        resource: str,
        export_format: str = "csv",
        filters: Optional[Dict[str, Any]] = None
    ) -> Tuple[io.BytesIO, str, str]:
        filters = filters or {}
        resource = resource.lower().strip()
        export_format = export_format.lower().strip()

        if export_format not in ("csv", "xlsx"):
            raise ValueError("Invalid export format. Supported formats: 'csv', 'xlsx'.")

        if resource == "applications":
            headers, rows, default_filename = self._export_applications(filters)
        elif resource == "assessments":
            headers, rows, default_filename = self._export_assessments(filters)
        elif resource in ("ai-interviews", "ai_interviews"):
            headers, rows, default_filename = self._export_ai_interviews(filters)
        elif resource == "interviews":
            headers, rows, default_filename = self._export_interviews(filters)
        elif resource == "offers":
            headers, rows, default_filename = self._export_offers(filters)
        elif resource in ("audit-logs", "audit_logs", "audit"):
            headers, rows, default_filename = self._export_audit_logs(filters)
        else:
            raise ValueError(f"Unsupported export resource: '{resource}'. Supported: applications, assessments, ai-interviews, interviews, offers, audit-logs.")

        if len(rows) > MAX_EXPORT_LIMIT:
            raise ValueError(f"Export result exceeds the {MAX_EXPORT_LIMIT} row safety limit (found {len(rows)} records). Please apply filters to narrow your export.")

        timestamp_str = datetime.utcnow().strftime("%Y-%m-%d")
        filename = f"{default_filename}_{timestamp_str}.{export_format}"

        if export_format == "csv":
            buffer = self._generate_csv(headers, rows)
            media_type = "text/csv; charset=utf-8"
        else:
            buffer = self._generate_xlsx(headers, rows, sheet_name=default_filename[:30])
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

        return buffer, filename, media_type

    def _generate_csv(self, headers: List[str], rows: List[List[Any]]) -> io.BytesIO:
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)
        writer.writerow(headers)
        for row in rows:
            # Clean string conversions
            clean_row = ["" if v is None else str(v) for v in row]
            writer.writerow(clean_row)
        
        # Encode to UTF-8 with BOM for Excel compatibility
        bytes_output = io.BytesIO(output.getvalue().encode('utf-8-sig'))
        bytes_output.seek(0)
        return bytes_output

    def _generate_xlsx(self, headers: List[str], rows: List[List[Any]], sheet_name: str = "Export") -> io.BytesIO:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

        # Header formatting
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center")

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        # Data formatting
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for row_idx, row in enumerate(rows, start=2):
            ws.append([None if v is None else str(v) for v in row])
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = thin_border

        # Auto-adjust column widths
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    def _export_applications(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("applications").select("candidate_name, email, phone, department, position, source, application_date, current_status, skills, portfolio, github, linkedin").order("application_date", desc=True)
        
        if filters.get("department") and filters["department"].lower() != 'all':
            query = query.eq("department", filters["department"])
        if filters.get("position_id") and filters["position_id"].lower() != 'all':
            query = query.eq("position_id", filters["position_id"])
        if filters.get("status") and filters["status"].lower() != 'all':
            query = query.eq("current_status", filters["status"])
        if filters.get("date_from"):
            query = query.gte("application_date", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("application_date", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Candidate Name", "Email", "Phone", "Department", "Position",
            "Source", "Application Date", "Status", "Skills",
            "Portfolio", "GitHub", "LinkedIn"
        ]

        rows = []
        for r in rows_data:
            # Format skills
            skills_val = r.get("skills")
            if isinstance(skills_val, list):
                skills_str = ", ".join(str(s) for s in skills_val)
            elif isinstance(skills_val, dict):
                skills_str = ", ".join(f"{k}: {v}" for k, v in skills_val.items())
            else:
                skills_str = str(skills_val or "")

            rows.append([
                r.get("candidate_name"),
                r.get("email"),
                r.get("phone"),
                r.get("department"),
                r.get("position"),
                r.get("source"),
                r.get("application_date"),
                r.get("current_status"),
                skills_str,
                r.get("portfolio"),
                r.get("github"),
                r.get("linkedin"),
            ])

        return headers, rows, "applications"

    def _export_assessments(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("assessments").select("score, pass_threshold, result, started_at, completed_at, created_at, applications(candidate_name, email, position, department)").order("created_at", desc=True)
        
        if filters.get("date_from"):
            query = query.gte("created_at", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("created_at", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Candidate Name", "Email", "Position", "Department",
            "Score", "Pass Threshold", "Result", "Started At", "Completed At"
        ]

        rows = []
        for r in rows_data:
            app = r.get("applications") or {}
            if filters.get("department") and filters["department"].lower() != 'all' and app.get("department") != filters["department"]:
                continue
            rows.append([
                app.get("candidate_name"),
                app.get("email"),
                app.get("position"),
                app.get("department"),
                r.get("score"),
                r.get("pass_threshold"),
                r.get("result"),
                r.get("started_at"),
                r.get("completed_at")
            ])

        return headers, rows, "assessments"

    def _export_ai_interviews(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("ai_interviews").select("started_at, completed_at, created_at, applications(candidate_name, email, position, department, current_status)").order("created_at", desc=True)
        
        if filters.get("date_from"):
            query = query.gte("created_at", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("created_at", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Candidate Name", "Email", "Position", "Department",
            "Started At", "Completed At", "Application Status"
        ]

        rows = []
        for r in rows_data:
            app = r.get("applications") or {}
            if filters.get("department") and filters["department"].lower() != 'all' and app.get("department") != filters["department"]:
                continue
            rows.append([
                app.get("candidate_name"),
                app.get("email"),
                app.get("position"),
                app.get("department"),
                r.get("started_at"),
                r.get("completed_at"),
                app.get("current_status")
            ])

        return headers, rows, "ai_interviews"

    def _export_interviews(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("interviews").select("round_number, date, time, status, attendance, applications(candidate_name, email, position, department), users!interviews_interviewer_id_fkey(email), interview_evaluations(overall_score, decision)").order("date", desc=True)
        
        if filters.get("date_from"):
            query = query.gte("date", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("date", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Candidate Name", "Email", "Position", "Department",
            "Round", "Interviewer Email", "Date", "Time",
            "Status", "Attendance", "Overall Score", "Decision"
        ]

        rows = []
        for r in rows_data:
            app = r.get("applications") or {}
            if filters.get("department") and filters["department"].lower() != 'all' and app.get("department") != filters["department"]:
                continue
            
            interviewer = r.get("users") or {}
            eval_data = r.get("interview_evaluations")
            ev = eval_data[0] if isinstance(eval_data, list) and eval_data else (eval_data if isinstance(eval_data, dict) else {})

            rows.append([
                app.get("candidate_name"),
                app.get("email"),
                app.get("position"),
                app.get("department"),
                r.get("round_number"),
                interviewer.get("email"),
                r.get("date"),
                r.get("time"),
                r.get("status"),
                r.get("attendance"),
                ev.get("overall_score") if ev else None,
                ev.get("decision") if ev else None
            ])

        return headers, rows, "interviews"

    def _export_offers(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("offers").select("email, designation, department, joining_date, duration, stipend, offer_issue_date, offer_expiry_date, offer_status, email_status, response_date, applications(candidate_name)").order("created_at", desc=True)
        
        if filters.get("offer_status") and filters["offer_status"].lower() != 'all':
            query = query.eq("offer_status", filters["offer_status"])
        if filters.get("department") and filters["department"].lower() != 'all':
            query = query.eq("department", filters["department"])
        if filters.get("date_from"):
            query = query.gte("offer_issue_date", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("offer_issue_date", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Candidate Name", "Email", "Designation", "Department",
            "Joining Date", "Duration", "Stipend", "Issue Date",
            "Expiry Date", "Offer Status", "Email Status", "Response Date"
        ]

        rows = []
        for r in rows_data:
            app = r.get("applications") or {}
            rows.append([
                app.get("candidate_name"),
                r.get("email"),
                r.get("designation"),
                r.get("department"),
                r.get("joining_date"),
                r.get("duration"),
                r.get("stipend"),
                r.get("offer_issue_date"),
                r.get("offer_expiry_date"),
                r.get("offer_status"),
                r.get("email_status"),
                r.get("response_date")
            ])

        return headers, rows, "offers"

    def _export_audit_logs(self, filters: Dict[str, Any]) -> Tuple[List[str], List[List[Any]], str]:
        query = supabase.table("audit_logs").select("action, role, application_id, previous_status, new_status, timestamp, users(email)").order("timestamp", desc=True)
        
        if filters.get("action") and filters["action"].lower() != 'all':
            query = query.eq("action", filters["action"])
        if filters.get("role") and filters["role"].lower() != 'all':
            query = query.eq("role", filters["role"])
        if filters.get("date_from"):
            query = query.gte("timestamp", filters["date_from"])
        if filters.get("date_to"):
            query = query.lte("timestamp", filters["date_to"])

        resp = query.limit(MAX_EXPORT_LIMIT + 1).execute()
        rows_data = resp.data or []

        headers = [
            "Action", "User Email", "Role", "Application ID",
            "Previous Status", "New Status", "Timestamp"
        ]

        rows = []
        for r in rows_data:
            user = r.get("users") or {}
            rows.append([
                r.get("action"),
                user.get("email") if user else r.get("hr_user"),
                r.get("role"),
                r.get("application_id"),
                r.get("previous_status"),
                r.get("new_status"),
                r.get("timestamp")
            ])

        return headers, rows, "audit_logs"
